from calendar import month_abbr
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from account.models import EmployeeSimpleInvoice

TEMPLATE_PATH = (
    Path(settings.BASE_DIR) / "account" / "invoice_templates" / "simple-invoice.xlsx"
)

SHEET_NAME = "Invoice"


def format_decimal(value: Decimal) -> Decimal:
    return Decimal(value).quantize(Decimal("0.01"))


def format_work_period(
    invoice: EmployeeSimpleInvoice,
) -> str:
    month = month_abbr[invoice.invoice_month]

    if invoice.start_day == invoice.end_day:
        return (
            f"Work period: " f"{month} {invoice.start_day}, " f"{invoice.invoice_year}"
        )

    return (
        f"Work period: "
        f"{month} {invoice.start_day}"
        f"–{invoice.end_day}, "
        f"{invoice.invoice_year}"
    )


def generate_simple_invoice_xlsx(
    *,
    invoice: EmployeeSimpleInvoice,
    employee_data: dict,
    output_path: Path,
) -> Path:
    """
    Generates a new XLSX invoice file based on the existing Excel template.

    Important:
    - Does not modify the original template.
    - Saves the generated file to output_path.
    """
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook[SHEET_NAME]

    employee = invoice.employee

    sheet["C4"] = invoice.invoice_number

    sheet["G4"] = timezone.localdate(invoice.created_at)

    sheet["C6"] = (f"{employee.last_name_en} " f"{employee.first_name_en}").strip()

    sheet["C7"] = employee_data["street_address"]

    sheet["C8"] = f'{employee_data["city"]}, ' f'{employee_data["province"]}'

    sheet["C9"] = employee_data["postal_code"]

    sheet["C10"] = invoice.gst_account_number or ""

    sheet.merge_cells("B14:D14")

    sheet["B14"] = format_work_period(invoice)

    sheet["B14"].alignment = Alignment(
        horizontal="left",
        vertical="center",
        shrink_to_fit=True,
    )

    sheet["E14"] = format_decimal(invoice.hours)

    sheet["F14"] = format_decimal(invoice.rate)

    gst = format_decimal(invoice.gst)

    wsbc = format_decimal(invoice.wsbc)

    sheet["G27"] = f"=G26*{gst}%"

    sheet["G28"] = f"=G26*{wsbc}%"

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(output_path)
    workbook.close()

    return output_path
